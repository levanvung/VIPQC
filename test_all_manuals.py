import os
import glob
import openpyxl
from bom_extractor import extract_full_bom
from excel_exporter import export_bom_to_excel, export_batch_bom_to_excel

def test_all():
    print("=" * 60)
    print("STARTING TEST ON ALL 5 SAMPLE MANUALS")
    print("=" * 60)
    
    out_dir = "excel_results"
    os.makedirs(out_dir, exist_ok=True)
    
    files = sorted(glob.glob("sample_manuals/*.pdf"))
    assert len(files) == 5, f"Expected 5 sample files, found {len(files)}"
    
    all_batch_data = []
    
    for f in files:
        fname = os.path.basename(f)
        print(f"\n--- Testing {fname} ---")
        meta, flat, mat = extract_full_bom(f)
        
        print(f"  PWB Code: {meta['pwb_code']}")
        print(f"  Processes: {meta['processes']}")
        print(f"  Total Unique Items: {meta['total_items']}")
        print(f"  Total Flat BOM Rows: {len(flat)}")
        print(f"  Total Parts Count: {meta['total_parts_count']}")
        
        assert meta['total_items'] > 0, f"No items extracted from {fname}"
        assert len(flat) > 0, f"No flat rows created for {fname}"
        assert meta['total_parts_count'] > 0, f"Total parts count is 0 for {fname}"
        
        # Test individual excel export
        base = os.path.splitext(fname)[0]
        excel_path = os.path.join(out_dir, f"BOM_{base}.xlsx")
        export_bom_to_excel(meta, flat, mat, excel_path)
        
        assert os.path.exists(excel_path), f"Excel file not created: {excel_path}"
        wb = openpyxl.load_workbook(excel_path)
        print(f"  Generated Excel: {os.path.basename(excel_path)} ({os.path.getsize(excel_path)} bytes)")
        print(f"  Sheets: {wb.sheetnames}")
        assert "Summary & Info" in wb.sheetnames
        assert "Consolidated BOM" in wb.sheetnames
        
        all_batch_data.append((meta, flat, mat))

    # Test master batch export
    print("\n--- Testing Master Batch Export ---")
    master_path = os.path.join(out_dir, "Master_Consolidated_BOM.xlsx")
    export_batch_bom_to_excel(all_batch_data, master_path)
    assert os.path.exists(master_path), f"Master file not created: {master_path}"
    wb_master = openpyxl.load_workbook(master_path)
    print(f"  Generated Master Excel: {os.path.basename(master_path)} ({os.path.getsize(master_path)} bytes)")
    print(f"  Master Sheets: {wb_master.sheetnames}")
    
    print("\n" + "=" * 60)
    print("ALL TESTS PASSED WITH 100% SUCCESS!")
    print("=" * 60)

if __name__ == "__main__":
    test_all()
