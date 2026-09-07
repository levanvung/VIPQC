"""
Excel Exporter for Working Manual BOM Extractor
Creates professional, publication-ready Excel workbooks with:
- Summary & Info sheet
- Consolidated BOM sheet
- BOM Matrix View sheet
- Individual tab for each unique PROCESS CODE
- Batch master workbook support
"""

import os
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─── Color Palette ────────────────────────────────────────────────────────────
COLOR_HEADER_BG     = "0D1B2A"   # Deep midnight navy
COLOR_HEADER_FG     = "FFFFFF"
COLOR_SUB_BG        = "1B4F72"   # Medium blue
COLOR_ACCENT        = "1ABC9C"   # Teal accent line
COLOR_ZEBRA         = "F0F4F8"   # Very light blue-grey
COLOR_BORDER        = "C5CDD6"
COLOR_PROCESS_TABS = {           # Per-process-type header colours
    "CHA": "154360",  # Deep blue  – Chip / SMT
    "SCP": "1B5E20",  # Deep green – Single / DIP
    "RAD": "7B341E",  # Deep brown – Radial
    "CHP": "4A148C",  # Purple     – CHP process
}
FONT = "Segoe UI"


# ─── Helpers ──────────────────────────────────────────────────────────────────
def _header_style(bg=COLOR_HEADER_BG, fg=COLOR_HEADER_FG, bold=True, size=10, center=True):
    return dict(
        font=Font(name=FONT, size=size, bold=bold, color=fg),
        fill=PatternFill(start_color=bg, end_color=bg, fill_type="solid"),
        alignment=Alignment(horizontal="center" if center else "left",
                            vertical="center", wrap_text=True),
        border=Border(
            top=Side(style="thin", color=COLOR_BORDER),
            bottom=Side(style="medium", color=bg),
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
        ),
    )


def _cell_style(zebra=False, align="left", bold=False, fg=None):
    bg = COLOR_ZEBRA if zebra else "FFFFFF"
    return dict(
        font=Font(name=FONT, size=10, bold=bold, color=fg or "212529"),
        fill=PatternFill(start_color=bg, end_color=bg, fill_type="solid"),
        alignment=Alignment(horizontal=align, vertical="center"),
        border=Border(
            top=Side(style="thin", color=COLOR_BORDER),
            bottom=Side(style="thin", color=COLOR_BORDER),
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
        ),
    )


def _apply(cell, style: dict):
    for k, v in style.items():
        setattr(cell, k, v)


def _autofit(ws, min_w=10, max_w=48):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        best = max(
            (len(str(c.value or "").split("\n")[0]) for c in col),
            default=0
        )
        ws.column_dimensions[letter].width = max(min_w, min(best + 4, max_w))


def _safe_sheet_name(proc: str) -> str:
    """Make a valid Excel sheet name ≤ 31 chars."""
    return re.sub(r'[\:\\/\?\*\[\]]', '_', proc)[:31]


def _process_header_color(proc: str) -> str:
    prefix = proc[:3].upper()
    return COLOR_PROCESS_TABS.get(prefix, COLOR_SUB_BG)


def _write_process_sheet(wb, proc_code: str, records: list):
    """Creates one Excel sheet for a single process code."""
    sheet_name = _safe_sheet_name(proc_code)
    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True

    hdr_color = _process_header_color(proc_code)

    # ── Section banner ─────────────────────────────────────────
    ws.merge_cells("A1:H1")
    banner = ws["A1"]
    banner.value = f"PROCESS CODE: {proc_code}  │  Total items: {len(records)}"
    _apply(banner, _header_style(bg=hdr_color, size=11))
    ws.row_dimensions[1].height = 28

    # ── Accent line ─────────────────────────────────────────────
    ws.merge_cells("A2:H2")
    accent = ws["A2"]
    accent.fill = PatternFill(start_color=COLOR_ACCENT, end_color=COLOR_ACCENT, fill_type="solid")
    ws.row_dimensions[2].height = 3

    # ── Column headers ──────────────────────────────────────────
    headers = [
        ("Item No",                         "center", 8),
        ("Part No / SNM Code",              "left",   22),
        ("Internal S.N.",                   "center", 16),
        ("Rating / Specification",          "left",   30),
        ("Model Variant",                   "center", 18),
        ("Qty",                             "center", 6),
        ("Reference Designators (Locations)","left",  40),
        ("Alternate Parts",                 "left",   22),
    ]
    HEADER_ROW = 3
    ws.row_dimensions[HEADER_ROW].height = 26
    for ci, (hname, align, w) in enumerate(headers, 1):
        c = ws.cell(row=HEADER_ROW, column=ci, value=hname)
        _apply(c, _header_style(bg=COLOR_HEADER_BG, size=10))
        ws.column_dimensions[get_column_letter(ci)].width = w

    # ── Data rows ───────────────────────────────────────────────
    for ri, r in enumerate(records, start=HEADER_ROW + 1):
        z = (ri % 2 == 0)
        ws.row_dimensions[ri].height = 20
        vals = [
            r["Item_No"],
            r["Part_No"],
            r["Internal_SN"],
            r["Rating_Spec"],
            r["Model_Variant"],
            r["Qty"],
            r["Remarks_Locations"],
            r["Alternate_Parts"],
        ]
        aligns = [h[1] for h in headers]
        for ci, (val, al) in enumerate(zip(vals, aligns), 1):
            is_qty = (ci == 6)
            c = ws.cell(row=ri, column=ci, value=val)
            _apply(c, _cell_style(zebra=z, align=al, bold=is_qty))

    # ── AutoFilter, freeze ──────────────────────────────────────
    end_col = get_column_letter(len(headers))
    end_row = HEADER_ROW + len(records)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{end_col}{end_row}"
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    # ── Summary footer ──────────────────────────────────────────
    footer_row = end_row + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=5)
    fc = ws.cell(row=footer_row, column=1, value=f"Total Qty Assembled: {sum(r['Qty'] for r in records if r['Qty'])}")
    fc.font = Font(name=FONT, size=10, bold=True, color=hdr_color)
    fc.alignment = Alignment(horizontal="right", vertical="center")

    return ws


# ─── Single-file Excel export ─────────────────────────────────────────────────
def export_bom_to_excel(doc_metadata: dict, consolidated_records: list,
                        matrix_records: list, output_path: str) -> str:
    wb = openpyxl.Workbook()

    # ── SHEET 1: Summary & Info ───────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary & Info"

    ws_sum.merge_cells("A1:F1")
    t = ws_sum["A1"]
    t.value = "WORKING MANUAL BOM EXTRACTION REPORT"
    t.font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
    t.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 44

    # accent bar
    ws_sum.merge_cells("A2:F2")
    ws_sum["A2"].fill = PatternFill(start_color=COLOR_ACCENT, end_color=COLOR_ACCENT, fill_type="solid")
    ws_sum.row_dimensions[2].height = 4

    ws_sum.merge_cells("A3:F3")
    sub = ws_sum["A3"]
    sub.value = f"Generated: {datetime.now().strftime('%Y-%m-%d  %H:%M:%S')}   |   Antigravity BOM Extractor"
    sub.font = Font(name=FONT, size=9, italic=True, color="6C757D")
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[3].height = 18

    # doc meta
    meta_rows = [
        ("Source File",           doc_metadata.get("filename", "")),
        ("PWB Part Code",         doc_metadata.get("pwb_code", "N/A")),
        ("Board Model",           doc_metadata.get("model", "N/A")),
        ("Total Pages",           str(doc_metadata.get("total_pages", 1))),
        ("Process Codes Detected",", ".join(doc_metadata.get("processes", [])) or "N/A"),
        ("Total Unique BOM Items",str(doc_metadata.get("total_items", 0))),
        ("Total Parts Qty",       str(doc_metadata.get("total_parts_count", 0))),
    ]
    ws_sum.cell(row=5, column=1, value="DOCUMENT SPECIFICATIONS").font = \
        Font(name=FONT, size=11, bold=True, color=COLOR_HEADER_BG)
    ws_sum.row_dimensions[5].height = 22

    for ri, (lbl, val) in enumerate(meta_rows, start=6):
        ws_sum.row_dimensions[ri].height = 22
        lc = ws_sum.cell(row=ri, column=1, value=lbl)
        lc.font = Font(name=FONT, size=10, bold=True)
        lc.fill = PatternFill(start_color="EAF2FB", end_color="EAF2FB", fill_type="solid")
        lc.border = Border(
            top=Side(style="thin", color=COLOR_BORDER),
            bottom=Side(style="thin", color=COLOR_BORDER),
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
        )
        ws_sum.merge_cells(start_row=ri, start_column=2, end_row=ri, end_column=4)
        vc = ws_sum.cell(row=ri, column=2, value=val)
        vc.font = Font(name=FONT, size=10)
        vc.border = Border(
            top=Side(style="thin", color=COLOR_BORDER),
            bottom=Side(style="thin", color=COLOR_BORDER),
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
        )

    # breakdown table
    breakdown_row = len(meta_rows) + 8
    ws_sum.cell(row=breakdown_row, column=1, value="PROCESS CODE BREAKDOWN").font = \
        Font(name=FONT, size=11, bold=True, color=COLOR_HEADER_BG)

    bh_row = breakdown_row + 1
    bhdrs = ["Process Code", "Model Variant", "Item Count", "Total Assembled Qty"]
    ws_sum.row_dimensions[bh_row].height = 24
    for ci, h in enumerate(bhdrs, 1):
        c = ws_sum.cell(row=bh_row, column=ci, value=h)
        _apply(c, _header_style(bg=COLOR_HEADER_BG, size=10))

    variant_stats: dict = {}
    for r in consolidated_records:
        key = (r["Process"], r["Model_Variant"])
        variant_stats.setdefault(key, {"items": 0, "qty": 0})
        variant_stats[key]["items"] += 1
        variant_stats[key]["qty"] += r["Qty"]

    for ri_offset, ((proc, vname), stats) in enumerate(sorted(variant_stats.items()), start=1):
        row = bh_row + ri_offset
        z = (ri_offset % 2 == 1)
        ws_sum.row_dimensions[row].height = 20
        tab_color = _process_header_color(proc)
        for ci, val in enumerate([proc, vname, stats["items"], stats["qty"]], 1):
            c = ws_sum.cell(row=row, column=ci, value=val)
            _apply(c, _cell_style(zebra=z, align="center"))
            if ci == 1:
                c.font = Font(name=FONT, size=10, bold=True, color=tab_color)

    _autofit(ws_sum)

    # ── SHEET 2: Consolidated BOM ─────────────────────────────────────────────
    ws_flat = wb.create_sheet(title="Consolidated BOM")

    flat_headers = [
        ("Seq",           "center", 6),
        ("Item No",       "center", 8),
        ("Process Code",  "center", 18),
        ("Model Variant", "center", 16),
        ("PWB Part Code", "center", 16),
        ("Part No / SNM Code", "left", 24),
        ("Internal S.N.", "center", 16),
        ("Rating / Specification", "left", 30),
        ("Qty",           "center", 6),
        ("Reference Designators / Locations", "left", 42),
        ("Alternate Parts", "left", 24),
        ("Page",          "center", 6),
        ("Source File",   "left",   28),
    ]
    ws_flat.row_dimensions[1].height = 28
    for ci, (hname, al, w) in enumerate(flat_headers, 1):
        c = ws_flat.cell(row=1, column=ci, value=hname)
        _apply(c, _header_style(bg=COLOR_HEADER_BG, size=10))
        ws_flat.column_dimensions[get_column_letter(ci)].width = w

    for ri, r in enumerate(consolidated_records, start=2):
        z = (ri % 2 == 0)
        ws_flat.row_dimensions[ri].height = 20
        vals = [
            r["Seq"], r["Item_No"], r["Process"], r["Model_Variant"],
            r["PWB_Code"], r["Part_No"], r["Internal_SN"], r["Rating_Spec"],
            r["Qty"], r["Remarks_Locations"], r["Alternate_Parts"],
            r["Source_Page"], r["Filename"]
        ]
        aligns = [h[1] for h in flat_headers]
        for ci, (val, al) in enumerate(zip(vals, aligns), 1):
            c = ws_flat.cell(row=ri, column=ci, value=val)
            _apply(c, _cell_style(zebra=z, align=al))

    end_col = get_column_letter(len(flat_headers))
    ws_flat.auto_filter.ref = f"A1:{end_col}{len(consolidated_records)+1}"
    ws_flat.freeze_panes = "A2"

    # ── SHEET 3: BOM Matrix View ──────────────────────────────────────────────
    if matrix_records:
        ws_mat = wb.create_sheet(title="BOM Matrix View")
        base_cols = ["No", "Page", "Process", "PWB_Code", "Part_No",
                     "Internal_SN", "Rating"]
        var_q = sorted({k for r in matrix_records for k in r if k.startswith("Qty_")})
        var_r = sorted({k for r in matrix_records for k in r if k.startswith("Remarks_")})
        all_cols = base_cols + var_q + var_r + ["Alternate_Parts"]

        ws_mat.row_dimensions[1].height = 28
        for ci, h in enumerate(all_cols, 1):
            c = ws_mat.cell(row=1, column=ci, value=h.replace("_", " "))
            _apply(c, _header_style(bg=COLOR_SUB_BG, size=10))

        for ri, r in enumerate(matrix_records, start=2):
            z = (ri % 2 == 0)
            ws_mat.row_dimensions[ri].height = 20
            for ci, h in enumerate(all_cols, 1):
                val = r.get(h, "")
                al = "center" if any(k in h for k in ["No","Page","Qty","SN"]) else "left"
                c = ws_mat.cell(row=ri, column=ci, value=val)
                _apply(c, _cell_style(zebra=z, align=al))

        end_c = get_column_letter(len(all_cols))
        ws_mat.auto_filter.ref = f"A1:{end_c}{len(matrix_records)+1}"
        ws_mat.freeze_panes = "A2"
        _autofit(ws_mat)

    # ── SHEETS 4+: One tab per PROCESS CODE ───────────────────────────────────
    processes = doc_metadata.get("processes", [])
    for proc in processes:
        proc_recs = [r for r in consolidated_records if r["Process"] == proc]
        if not proc_recs:
            continue
        _write_process_sheet(wb, proc, proc_recs)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path


# ─── Batch (multi-file) master export ────────────────────────────────────────
def export_batch_bom_to_excel(all_results: list, output_path: str) -> str:
    """
    all_results: list of (metadata, consolidated_records, matrix_records)
    """
    wb = openpyxl.Workbook()

    # ── SHEET 1: Batch Summary ────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Batch Summary"

    ws_sum.merge_cells("A1:G1")
    t = ws_sum["A1"]
    t.value = "BATCH WORKING MANUAL BOM — MASTER CONSOLIDATED REPORT"
    t.font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    t.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 40

    ws_sum.merge_cells("A2:G2")
    ws_sum["A2"].fill = PatternFill(start_color=COLOR_ACCENT, end_color=COLOR_ACCENT, fill_type="solid")
    ws_sum.row_dimensions[2].height = 4

    ws_sum.merge_cells("A3:G3")
    s = ws_sum["A3"]
    s.value = (f"Manuals: {len(all_results)}  |  "
               f"Generated: {datetime.now().strftime('%Y-%m-%d  %H:%M:%S')}  |  Antigravity BOM Extractor")
    s.font = Font(name=FONT, size=9, italic=True, color="6C757D")
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[3].height = 18

    bhdrs = ["No", "Manual File", "PWB Code", "Board Model",
             "Process Codes", "BOM Items", "Total Qty"]
    ws_sum.row_dimensions[5].height = 26
    for ci, h in enumerate(bhdrs, 1):
        c = ws_sum.cell(row=5, column=ci, value=h)
        _apply(c, _header_style(bg=COLOR_HEADER_BG, size=10))

    for idx, (meta, flat, _) in enumerate(all_results, 1):
        ri = idx + 5
        z = (idx % 2 == 1)
        ws_sum.row_dimensions[ri].height = 20
        vals = [
            idx,
            meta.get("filename", ""),
            meta.get("pwb_code", "N/A"),
            meta.get("model", "N/A"),
            ", ".join(meta.get("processes", [])),
            meta.get("total_items", 0),
            meta.get("total_parts_count", 0),
        ]
        aligns = ["center","left","center","left","left","center","center"]
        for ci, (val, al) in enumerate(zip(vals, aligns), 1):
            c = ws_sum.cell(row=ri, column=ci, value=val)
            _apply(c, _cell_style(zebra=z, align=al))

    _autofit(ws_sum)

    # ── SHEET 2: All Manuals Consolidated BOM ─────────────────────────────────
    ws_all = wb.create_sheet(title="All Manuals — Consolidated BOM")
    all_hdrs = [
        ("Seq",           "center", 6),
        ("Manual File",   "left",   26),
        ("Item No",       "center", 8),
        ("Process Code",  "center", 20),
        ("Model Variant", "center", 18),
        ("PWB Part Code", "center", 16),
        ("Part No / SNM", "left",   24),
        ("Internal S.N.", "center", 16),
        ("Rating / Spec", "left",   30),
        ("Qty",           "center", 6),
        ("Designators",   "left",   42),
        ("Alternate",     "left",   22),
        ("Page",          "center", 6),
    ]
    ws_all.row_dimensions[1].height = 28
    for ci, (h, al, w) in enumerate(all_hdrs, 1):
        c = ws_all.cell(row=1, column=ci, value=h)
        _apply(c, _header_style(bg=COLOR_HEADER_BG, size=10))
        ws_all.column_dimensions[get_column_letter(ci)].width = w

    seq = 1
    ri = 2
    for meta, flat, _ in all_results:
        for r in flat:
            z = (ri % 2 == 0)
            ws_all.row_dimensions[ri].height = 20
            vals = [
                seq, r["Filename"], r["Item_No"], r["Process"],
                r["Model_Variant"], r["PWB_Code"], r["Part_No"],
                r["Internal_SN"], r["Rating_Spec"], r["Qty"],
                r["Remarks_Locations"], r["Alternate_Parts"], r["Source_Page"]
            ]
            aligns = [h[1] for h in all_hdrs]
            for ci, (val, al) in enumerate(zip(vals, aligns), 1):
                c = ws_all.cell(row=ri, column=ci, value=val)
                _apply(c, _cell_style(zebra=z, align=al))
            seq += 1
            ri += 1

    ws_all.auto_filter.ref = f"A1:{get_column_letter(len(all_hdrs))}{ri-1}"
    ws_all.freeze_panes = "A2"

    # ── SHEETS 3+: One tab per PROCESS CODE (all manuals combined) ────────────
    # Gather all unique process codes across all manuals
    all_processes: dict = {}   # process_code -> list of records
    for meta, flat, _ in all_results:
        for r in flat:
            proc = r["Process"]
            all_processes.setdefault(proc, []).append(r)

    for proc_code in sorted(all_processes.keys()):
        _write_process_sheet(wb, proc_code, all_processes[proc_code])

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path
