"""
BOM Comparator Engine
Provides bidirectional comparison between an engineering/production Excel BOM
and an extracted Working Manual (PDF) BOM page.
Detects:
  - Exact matches (Part No, Qty, Reference Designators)
  - Location/Designator discrepancies (missing in WM or missing in Excel)
  - Quantity discrepancies
  - Missing parts in WM or missing parts in Excel
Also exports color-coded comparison audit reports to Excel.
"""

import os
import sys
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure Windows C-runtime DLLs are loaded for pymupdf
if sys.platform == "win32":
    for p in [r"C:\Program Files\Common Files\microsoft shared\ClickToRun", r"C:\Program Files\Microsoft Office\root\Client"]:
        if os.path.exists(p):
            try:
                os.add_dll_directory(p)
            except Exception:
                pass

import pymupdf
from bom_extractor import parse_page_bom, sanitize_part_and_rating


def normalize_part_no(pn: str) -> str:
    """Normalizes part numbers for reliable cross-document matching."""
    if not pn:
        return ""
    # Remove leading/trailing spaces, trailing '+' or '-', uppercase
    cleaned = re.sub(r'[\s\+\-]+$', '', str(pn).strip()).upper()
    # Also strip spaces around hyphens or slashes
    cleaned = re.sub(r'\s*([-\/])\s*', r'\1', cleaned)
    return cleaned


def normalize_key(pn: str) -> str:
    """Strict alphanumeric key for fuzzy comparison when symbols differ slightly."""
    return re.sub(r'[^A-Z0-9]', '', normalize_part_no(pn))


# Electronic Reference Designator pattern: 1-4 letters + 1+ digits + optional suffix (e.g. C1, R10, D922, IC1, CN101, C921-1, U1.1)
REF_DES_PATTERN = re.compile(r'^[A-Za-z]{1,4}\d+[A-Za-z0-9\-_\.]*$')

# Max allowed file size (10 MB) to prevent lag and ensure responsive performance
MAX_FILE_SIZE_MB = 10
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

# Metadata / approval / non-BOM keywords to filter out footer rows and non-component data
NON_BOM_KEYWORDS = [
    # Chinese approval / metadata / statistics
    "编制", "审核", "批准", "确认", "核准", "制表", "制单", "担当", "承认", 
    "单板点数", "总点数", "点数", "标准工时", "工时", "更新日期", "BOM更新日期",
    "合计", "小计", "总计", "备注", "版本",
    # English keywords
    "REMARK", "REMARKS", "NOTE", "NOTES", "FEEDING DIRECTION", "DIRECTION",
    "PREPARED", "CHECKED", "APPROVED", "CONFIRMED", "SIGNATURE", "AUDIT",
    "TOTAL", "SUBTOTAL", "POINTS", "CYCLE TIME", "TACT TIME", "STD TIME",
    "REVISION", "REV DATE", "DATE:",
    # Vietnamese keywords
    "NGƯỜI LẬP", "KIỂM TRA", "PHÊ DUYỆT", "XÁC NHẬN", "KÝ TÊN", 
    "TỔNG CỘNG", "GHI CHÚ", "LƯU Ý", "CHÚ Ý"
]


def is_valid_part_no(pn: str) -> bool:
    """Validates if a string looks like an authentic engineering/electronic Part Number."""
    if not pn:
        return False
    pn_clean = str(pn).strip()
    if len(pn_clean) < 2 or len(pn_clean) > 40:
        return False
    pn_upper = pn_clean.upper()
    if any(k.upper() in pn_upper for k in NON_BOM_KEYWORDS):
        return False
    if ":" in pn_clean or "：" in pn_clean:
        return False
    if len(pn_clean.split()) > 3:
        return False
    if not re.search(r'[A-Za-z0-9]', pn_clean):
        return False
    return True


def parse_locations_string(raw_locs: str) -> list[str]:
    """Splits raw location text into sorted, deduplicated reference designators."""
    if not raw_locs:
        return []
    # Split by spaces, commas, semicolons, tabs, newlines
    tokens = re.split(r'[,;\s\n\r\t]+', str(raw_locs).strip())
    cleaned = [t.strip().upper() for t in tokens if REF_DES_PATTERN.match(t.strip())]
    return sorted(list(dict.fromkeys(cleaned)))


def parse_excel_bom(filepath: str, sheet_name: str = None) -> dict:
    """
    Parses an Excel BOM file (supports SMT/AI recipes, factory BOMs, or app exports).
    Automatically identifies header row and semantic columns.
    Returns:
      - filename, filepath, sheet_name, all_sheets
      - process_code (if found in metadata cells)
      - title (if found in header cells)
      - items: list of parsed item dicts
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Excel file not found: {filepath}")
    if os.path.getsize(filepath) > MAX_FILE_SIZE_BYTES:
        raise ValueError(f"Dung lượng file '{os.path.basename(filepath)}' vượt quá giới hạn tối đa ({MAX_FILE_SIZE_MB}MB).")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheets = wb.sheetnames
    target_sheet = sheet_name if sheet_name and sheet_name in sheets else sheets[0]
    ws = wb[target_sheet]

    # Inspect top 25 rows to find metadata and header row
    header_row_idx = None
    col_mapping = {}
    detected_process = ""
    detected_title = ""

    # Column classification rules
    rules = {
        "station": ["站位", "STATION", "STN", "ITEM", "ITEM NO", "NO", "NO.", "SEQ", "STT"],
        "part_no": ["物料编码", "PART NO", "PART NO.", "PART NUMBER", "PART", "MÃ LINH KIỆN", "MÃ LK", "CODE", "料号", "PART_NO"],
        "rating": ["规格", "SPEC", "SPECIFICATION", "RATING", "RATING / VALUE", "VALUE", "QUY CÁCH"],
        "polarity": ["极性", "POLARITY", "POL", "CỰC TÍNH"],
        "locations": ["位置", "LOCATION", "LOCATIONS", "REMARKS", "DESIGNATORS", "REF DES", "REFERENCE DESIGNATORS", "VỊ TRÍ"],
        "qty": ["数量", "QTY", "QTY.", "QUANTITY", "SỐ LƯỢNG"]
    }

    max_search_row = min(25, ws.max_row)
    for r in range(1, max_search_row + 1):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, min(ws.max_column + 1, 30))]

        # Check for process metadata (e.g. 程序: SCP3072AR-2G)
        for c_idx, val in enumerate(row_vals):
            val_u = val.upper()
            if any(k in val_u for k in ["程序", "PROCESS", "PROGRAM", "PROCESS CODE"]):
                m = re.search(r'(?:程序|PROCESS|PROGRAM|PROCESS CODE)[\s:：]+([A-Z0-9\-_]+)', val, re.I)
                if m:
                    detected_process = m.group(1).strip()
                elif c_idx + 1 < len(row_vals) and row_vals[c_idx + 1]:
                    detected_process = row_vals[c_idx + 1].strip()

            # Board model / title check
            if r == 1 and val and len(val) >= 3 and not detected_title and not any(k in val_u for k in ["程序", "PROCESS", "BOM"]):
                detected_title = val

        # Check if this row is the column header row
        matched_roles = {}
        for c_idx, val in enumerate(row_vals):
            v_clean = val.upper()
            if not v_clean:
                continue
            for role, keywords in rules.items():
                if role not in matched_roles:
                    if any(v_clean == k or v_clean.startswith(k) or f" {k} " in f" {v_clean} " for k in keywords):
                        matched_roles[role] = c_idx + 1
                        break

        # A valid BOM header row must have at least part_no or (locations and rating)
        if "part_no" in matched_roles or ("locations" in matched_roles and "rating" in matched_roles):
            header_row_idx = r
            col_mapping = matched_roles
            break

    # If header not found by keyword, fallback to default positions if rows exist
    if not header_row_idx:
        header_row_idx = 3 if ws.max_row >= 3 else 1
        col_mapping = {
            "station": 1,
            "part_no": 2,
            "rating": 3,
            "polarity": 4,
            "locations": 5
        }

    # Parse rows
    parsed_items = []
    item_seq = 1

    for r in range(header_row_idx + 1, ws.max_row + 1):
        def get_val(role):
            col = col_mapping.get(role)
            if col:
                v = ws.cell(row=r, column=col).value
                return str(v).strip() if v is not None else ""
            return ""

        # Check entire row text for metadata / non-BOM keywords (footers, sign-offs, notes)
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, min(ws.max_column + 1, 30))]
        row_text = " ".join(v for v in row_vals if v)
        row_upper = row_text.upper()

        if not row_upper:
            continue

        if any(k.upper() in row_upper for k in NON_BOM_KEYWORDS):
            continue

        part_raw = get_val("part_no")
        locs_raw = get_val("locations")
        rating_raw = get_val("rating")
        stn_raw = get_val("station")
        pol_raw = get_val("polarity")
        qty_raw = get_val("qty")

        # Skip completely empty rows
        if not part_raw and not locs_raw and not rating_raw:
            continue

        # Skip sub-headers or summary rows
        if any(k in part_raw.upper() for k in ["TOTAL", "TỔNG", "合计", "小计", "PAGE", "NOTE", "REMARK", "点数"]):
            continue

        locations = parse_locations_string(locs_raw)
        valid_pn = is_valid_part_no(part_raw)

        # Fallback if part number was placed in station column and station column has valid PN
        if not valid_pn and is_valid_part_no(stn_raw) and not locations:
            part_raw = stn_raw
            valid_pn = True

        # If neither a valid part number nor valid locations exist, this is a non-BOM row
        if not valid_pn and not locations:
            continue

        # Quantity calculation
        qty = 0
        if qty_raw:
            try:
                qty = int(float(qty_raw))
            except ValueError:
                qty = len(locations)
        else:
            qty = len(locations)

        parsed_items.append({
            "seq": item_seq,
            "row_index": r,
            "station": stn_raw or str(item_seq),
            "part_no": part_raw,
            "norm_part_no": normalize_part_no(part_raw),
            "norm_key": normalize_key(part_raw),
            "rating": rating_raw,
            "polarity": pol_raw,
            "qty": qty,
            "locations_raw": locs_raw,
            "locations": locations
        })
        item_seq += 1

    wb.close()

    return {
        "filepath": filepath,
        "filename": os.path.basename(filepath),
        "sheet_name": target_sheet,
        "all_sheets": sheets,
        "process_code": detected_process,
        "title": detected_title,
        "header_row": header_row_idx,
        "total_items": len(parsed_items),
        "items": parsed_items
    }


def get_wm_pages_info(pdf_path: str) -> list[dict]:
    """
    Extracts summary information for each page in a Working Manual PDF.
    Returns a list of dicts with page index, page number, process code, variants, and item count.
    """
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")
    if os.path.getsize(pdf_path) > MAX_FILE_SIZE_BYTES:
        raise ValueError(f"Dung lượng file '{os.path.basename(pdf_path)}' vượt quá giới hạn tối đa ({MAX_FILE_SIZE_MB}MB).")

    doc = pymupdf.open(pdf_path)
    total_pages = len(doc)
    pages_info = []

    for pno in range(total_pages):
        try:
            meta, items = parse_page_bom(doc, pno)
            pwb = meta.get("pwb_code", "")
            proc = meta.get("process_code", "")

            # Collect variants from items
            variants = []
            for it in items:
                vkeys = list(it.get("variants_data", {}).keys())
                if vkeys:
                    variants = vkeys
                    break

            # Count items having at least one variant with qty > 0 or remarks
            active_count = 0
            for it in items:
                vdata = it.get("variants_data", {})
                if any(vd.get("qty", 0) > 0 or vd.get("remarks") for vd in vdata.values()):
                    active_count += 1
                elif it.get("part_no") or it.get("sn"):
                    active_count += 1

            effective_proc = proc or (variants[0] if variants else (pwb or f"Trang {pno + 1}"))
            var_desc = f" [{', '.join(variants)}]" if variants else ""

            pages_info.append({
                "page_index": pno,
                "page_num": pno + 1,
                "total_pages": total_pages,
                "pwb_code": pwb,
                "process_code": effective_proc,
                "variants": variants,
                "items_count": active_count,
                "display_text": f"Trang {pno + 1}/{total_pages}: {effective_proc}{var_desc} ({active_count} linh kiện)"
            })
        except Exception as e:
            pages_info.append({
                "page_index": pno,
                "page_num": pno + 1,
                "total_pages": total_pages,
                "pwb_code": "",
                "process_code": f"Trang {pno + 1}",
                "variants": [],
                "items_count": 0,
                "display_text": f"Trang {pno + 1}/{total_pages}: (Lỗi đọc: {e})"
            })

    doc.close()
    return pages_info


def extract_wm_page_items(pdf_path: str, page_index: int, variant_name: str = None) -> dict:
    """
    Extracts structured BOM items for a specific page and variant of a Working Manual PDF.
    """
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")
    if os.path.getsize(pdf_path) > MAX_FILE_SIZE_BYTES:
        raise ValueError(f"Dung lượng file '{os.path.basename(pdf_path)}' vượt quá giới hạn tối đa ({MAX_FILE_SIZE_MB}MB).")

    doc = pymupdf.open(pdf_path)
    if page_index < 0 or page_index >= len(doc):
        doc.close()
        raise IndexError(f"Page index {page_index} out of range (0-{len(doc)-1})")

    meta, raw_items = parse_page_bom(doc, page_index)
    doc.close()

    # Collect variants from items
    variants = []
    for it in raw_items:
        vkeys = list(it.get("variants_data", {}).keys())
        if vkeys:
            variants = vkeys
            break

    # Pick active variant (if user specified or default to the one with most active items)
    if variant_name and variant_name in variants:
        active_variant = variant_name
    elif variants:
        # Check which variant has more items with qty > 0
        best_v = variants[0]
        max_q = -1
        for v in variants:
            q_sum = sum(it.get("variants_data", {}).get(v, {}).get("qty", 0) for it in raw_items)
            rem_cnt = sum(len(it.get("variants_data", {}).get(v, {}).get("remarks", [])) for it in raw_items)
            score = q_sum + rem_cnt
            if score > max_q:
                max_q = score
                best_v = v
        active_variant = best_v
    else:
        active_variant = ""

    extracted_items = []
    item_seq = 1

    for it in raw_items:
        raw_pn = it.get("part_no", "")
        raw_sn = it.get("sn", "")
        raw_rating = it.get("rating", "")

        part_no, rating = sanitize_part_and_rating(raw_pn, raw_rating)
        effective_pn = part_no or raw_sn

        vdata = it.get("variants_data", {}).get(active_variant, {})
        raw_remarks = vdata.get("remarks", [])
        qty = vdata.get("qty", 0)

        clean_locs = parse_locations_string(" ".join(raw_remarks))
        effective_qty = len(clean_locs) if clean_locs else qty

        if effective_pn or clean_locs or effective_qty > 0:
            extracted_items.append({
                "seq": item_seq,
                "item_no": it.get("item_no", item_seq),
                "part_no": effective_pn,
                "raw_part_no": raw_pn,
                "sn": raw_sn,
                "norm_part_no": normalize_part_no(effective_pn),
                "norm_key": normalize_key(effective_pn),
                "rating": rating,
                "qty": effective_qty,
                "locations": clean_locs,
                "locations_raw": " ".join(clean_locs),
                "alternates": it.get("alternates", [])
            })
            item_seq += 1

    proc = meta.get("process_code", "") or (active_variant or meta.get("pwb_code", f"Trang {page_index + 1}"))
    return {
        "page_index": page_index,
        "page_num": page_index + 1,
        "pwb_code": meta.get("pwb_code", ""),
        "process_code": proc,
        "active_variant": active_variant,
        "variants": variants,
        "total_items": len(extracted_items),
        "items": extracted_items
    }


def compare_boms(excel_data: dict, wm_page_data: dict) -> dict:
    """
    Performs full bidirectional comparison between Excel BOM items and Working Manual page items.
    Returns:
      - is_all_matched: bool
      - total_excel, total_wm
      - count_matched, count_mismatched, count_missing_wm, count_missing_excel
      - comparison_rows: list of formatted comparison row records
    """
    excel_items = excel_data.get("items", [])
    wm_items = wm_page_data.get("items", [])

    wm_by_norm = {}
    wm_by_key = {}
    for idx, wi in enumerate(wm_items):
        norm = wi["norm_part_no"]
        key = wi["norm_key"]
        wm_by_norm.setdefault(norm, []).append((idx, wi))
        if key:
            wm_by_key.setdefault(key, []).append((idx, wi))

    used_wm_indices = set()
    comparison_rows = []

    for ei in excel_items:
        e_norm = ei["norm_part_no"]
        e_key = ei["norm_key"]
        e_locs = set(ei["locations"])
        e_qty = ei["qty"]

        candidates = wm_by_norm.get(e_norm, [])
        if not candidates and e_key:
            candidates = wm_by_key.get(e_key, [])

        avail_cands = [(w_idx, wi) for w_idx, wi in candidates if w_idx not in used_wm_indices]

        best_candidate = None
        matched_wm_indices = []

        # 1. Exact single match: Check if any single candidate matches e_locs exactly
        for w_idx, wi in avail_cands:
            if set(wi["locations"]) == e_locs and len(e_locs) > 0:
                best_candidate = wi
                matched_wm_indices = [w_idx]
                break

        # 2. Multi-row candidate aggregation: In SMT/AI manuals, a single BOM part number
        # is often split across multiple stations/feeders. If multiple candidate rows exist,
        # combine those whose locations intersect with e_locs or form the full set of locations.
        if not best_candidate and len(avail_cands) > 1:
            cands_with_overlap = [(w_idx, wi) for w_idx, wi in avail_cands if (set(wi["locations"]) & e_locs)]
            if cands_with_overlap:
                combined_locs = set().union(*(set(wi["locations"]) for _, wi in cands_with_overlap))
                if combined_locs == e_locs or len(cands_with_overlap) > 1:
                    combined_qty = sum(wi["qty"] for _, wi in cands_with_overlap)
                    primary_wi = cands_with_overlap[0][1]
                    merged_wi = dict(primary_wi)
                    merged_wi["qty"] = combined_qty
                    merged_wi["locations"] = sorted(list(combined_locs))
                    merged_wi["locations_raw"] = " ".join(sorted(list(combined_locs)))
                    best_candidate = merged_wi
                    matched_wm_indices = [w_idx for w_idx, _ in cands_with_overlap]

        # 3. Best single candidate by overlap
        if not best_candidate and avail_cands:
            best_wi = None
            best_idx = None
            best_overlap = -1
            for w_idx, wi in avail_cands:
                overlap = len(e_locs & set(wi["locations"]))
                if overlap > best_overlap:
                    best_overlap = overlap
                    best_wi = wi
                    best_idx = w_idx
            if best_wi is not None:
                best_candidate = best_wi
                matched_wm_indices = [best_idx]

        # 4. Fallback if no available candidate left but candidates existed in WM
        if not best_candidate and candidates:
            best_candidate = candidates[0][1]
            matched_wm_indices = []

        if best_candidate is not None:
            for idx in matched_wm_indices:
                used_wm_indices.add(idx)
            w_locs = set(best_candidate["locations"])
            w_qty = best_candidate["qty"]

            loc_missing_in_wm = sorted(list(e_locs - w_locs))
            loc_extra_in_wm = sorted(list(w_locs - e_locs))

            is_loc_match = (e_locs == w_locs)
            is_qty_match = (e_qty == w_qty)

            issues = []
            if not is_loc_match:
                if loc_missing_in_wm:
                    issues.append(f"WM thiếu vị trí: {', '.join(loc_missing_in_wm)}")
                if loc_extra_in_wm:
                    issues.append(f"Excel thiếu vị trí: {', '.join(loc_extra_in_wm)}")

            if not is_qty_match:
                issues.append(f"Lệch số lượng (Excel: {e_qty} ≠ WM: {w_qty})")

            rating_diff = False
            e_rat_clean = re.sub(r'\s+', ' ', ei["rating"]).strip()
            w_rat_clean = re.sub(r'\s+', ' ', best_candidate["rating"]).strip()
            if e_rat_clean and w_rat_clean and e_rat_clean.upper() != w_rat_clean.upper():
                rating_diff = True

            if is_loc_match and is_qty_match:
                status = "MATCH"
                status_label = "✅ KHỚP"
                if len(matched_wm_indices) > 1:
                    diff_summary = f"Khớp hoàn toàn (Gộp {len(matched_wm_indices)} dòng trạm trong WM)"
                else:
                    diff_summary = "Khớp hoàn toàn" if not rating_diff else "Khớp linh kiện & vị trí (Khác chuỗi quy cách)"
            else:
                status = "MISMATCH"
                status_label = "⚠️ SAI LỆCH"
                diff_summary = " | ".join(issues)

            comparison_rows.append({
                "status": status,
                "status_label": status_label,
                "part_no_excel": ei["part_no"],
                "part_no_wm": best_candidate["part_no"],
                "station": ei["station"],
                "qty_excel": e_qty,
                "qty_wm": w_qty,
                "locations_excel": " ".join(ei["locations"]),
                "locations_wm": " ".join(best_candidate["locations"]),
                "loc_missing_in_wm": loc_missing_in_wm,
                "loc_extra_in_wm": loc_extra_in_wm,
                "rating_excel": ei["rating"],
                "rating_wm": best_candidate["rating"],
                "diff_summary": diff_summary,
                "origin": "EXCEL"
            })
        else:
            comparison_rows.append({
                "status": "MISSING_IN_WM",
                "status_label": "❌ THIẾU Ở WM",
                "part_no_excel": ei["part_no"],
                "part_no_wm": "-",
                "station": ei["station"],
                "qty_excel": e_qty,
                "qty_wm": 0,
                "locations_excel": " ".join(ei["locations"]),
                "locations_wm": "-",
                "loc_missing_in_wm": ei["locations"],
                "loc_extra_in_wm": [],
                "rating_excel": ei["rating"],
                "rating_wm": "-",
                "diff_summary": "Không tìm thấy mã linh kiện trên trang WM này",
                "origin": "EXCEL"
            })

    for idx, wi in enumerate(wm_items):
        if idx not in used_wm_indices:
            w_locs_str = " ".join(wi["locations"]) if wi["locations"] else "(Không có vị trí)"
            comparison_rows.append({
                "status": "MISSING_IN_EXCEL",
                "status_label": "⚠️ THIẾU Ở EXCEL",
                "part_no_excel": "(File Excel thiếu)",
                "part_no_wm": wi["part_no"],
                "station": "-",
                "qty_excel": 0,
                "qty_wm": wi["qty"],
                "locations_excel": "(Chưa có)",
                "locations_wm": " ".join(wi["locations"]),
                "loc_missing_in_wm": [],
                "loc_extra_in_wm": wi["locations"],
                "rating_excel": "-",
                "rating_wm": wi["rating"],
                "diff_summary": f"File Excel đang thiếu linh kiện này! (WM có {wi['qty']} con tại vị trí: {w_locs_str})",
                "origin": "WM"
            })

    count_matched = sum(1 for r in comparison_rows if r["status"] == "MATCH")
    count_mismatched = sum(1 for r in comparison_rows if r["status"] == "MISMATCH")
    count_missing_wm = sum(1 for r in comparison_rows if r["status"] == "MISSING_IN_WM")
    count_missing_excel = sum(1 for r in comparison_rows if r["status"] == "MISSING_IN_EXCEL")
    
    # Total discrepancies includes all differences: Mismatches, Missing in WM, and Missing in Excel
    total_discrepancies = count_mismatched + count_missing_wm + count_missing_excel
    is_all_matched = (total_discrepancies == 0 and len(excel_items) > 0)

    return {
        "is_all_matched": is_all_matched,
        "total_excel": len(excel_items),
        "total_wm": len(wm_items),
        "total_rows": len(comparison_rows),
        "count_matched": count_matched,
        "count_mismatched": count_mismatched,
        "count_missing_wm": count_missing_wm,
        "count_missing_excel": count_missing_excel,
        "total_discrepancies": total_discrepancies,
        "excel_file": excel_data.get("filename", ""),
        "excel_process": excel_data.get("process_code", ""),
        "wm_process": wm_page_data.get("process_code", ""),
        "wm_page": wm_page_data.get("page_num", 1),
        "wm_variant": wm_page_data.get("active_variant", ""),
        "comparison_rows": comparison_rows
    }


def export_comparison_excel(comp_result: dict, output_path: str) -> str:
    """
    Exports a professional, beautifully styled Excel comparison audit report.
    Highlight colors:
      - Soft Green: Matched rows
      - Soft Amber/Orange: Location or Qty discrepancies
      - Soft Coral/Red: Missing in WM or Missing in Excel
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM Comparison Report"

    font_hero = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
    font_sub = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    font_sec = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_th = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=9, color="1E293B")
    font_mono = Font(name="Consolas", size=9, color="0F172A")
    font_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")

    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_match = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_mismatch = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_missing = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    ws.append(["BÁO CÁO ĐỐI CHIẾU SO SÁNH BOM (BOM COMPARISON AUDIT REPORT)"])
    ws.cell(row=1, column=1).font = font_hero
    ws.append([f"Thời gian xuất: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  Tạo bởi: Antigravity BOM Extractor & Comparator"])
    ws.cell(row=2, column=1).font = font_sub
    ws.append([])

    ws.append(["THÔNG TIN ĐỐI CHIẾU", "", "KẾT QUẢ TỔNG QUAN"])
    ws.cell(row=4, column=1).font = font_sec
    ws.cell(row=4, column=3).font = font_sec

    verdict = "✅ HOÀN TOÀN TRÙNG KHỚP (100% MATCHED)" if comp_result["is_all_matched"] else f"⚠️ PHÁT HIỆN {comp_result['total_discrepancies']} ĐIỂM SAI LỆCH"
    summary_data = [
        ("File Excel", comp_result["excel_file"], "Tổng số mục Excel", comp_result["total_excel"]),
        ("Quy trình Excel", comp_result["excel_process"] or "N/A", "Tổng số mục WM", comp_result["total_wm"]),
        ("Trang WM", f"Trang {comp_result['wm_page']} ({comp_result['wm_process']})", "Số mục khớp 100%", comp_result["count_matched"]),
        ("Model Variant", comp_result["wm_variant"] or "Mặc định", "Lệch vị trí / SL", comp_result["count_mismatched"]),
        ("Kết quả chung", verdict, "Thiếu trong WM / Excel", comp_result["count_missing_wm"] + comp_result["count_missing_excel"]),
    ]

    for label_l, val_l, label_r, val_r in summary_data:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=label_l).font = font_bold
        ws.cell(row=r, column=2, value=val_l).font = font_data
        ws.cell(row=r, column=3, value=label_r).font = font_bold
        c_r = ws.cell(row=r, column=4, value=val_r)
        c_r.font = font_bold
        if label_l == "Kết quả chung":
            c_r.font = Font(name="Segoe UI", size=10, bold=True, color="15803D" if comp_result["is_all_matched"] else "B91C1C")

    ws.append([])

    headers = [
        "STT",
        "Trạng Thái",
        "Mã LK (Excel)",
        "Mã LK (WM)",
        "SL (Excel)",
        "SL (WM)",
        "Vị Trí Cắm (Excel)",
        "Vị Trí Cắm (WM)",
        "Quy Cách / Spec (Excel)",
        "Quy Cách / Spec (WM)",
        "Chi Tiết Sai Lệch / Ghi Chú"
    ]

    h_row = ws.max_row + 1
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=h_row, column=col_idx)
        cell.font = font_th
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 5, 6] else "left", vertical="center")

    for idx, row in enumerate(comp_result["comparison_rows"], 1):
        status = row["status"]
        row_vals = [
            idx,
            row["status_label"],
            row["part_no_excel"],
            row["part_no_wm"],
            row["qty_excel"],
            row["qty_wm"],
            row["locations_excel"],
            row["locations_wm"],
            row["rating_excel"],
            row["rating_wm"],
            row["diff_summary"]
        ]
        r_num = ws.max_row + 1
        ws.append(row_vals)

        if status == "MATCH":
            row_fill = fill_match
        elif status == "MISMATCH":
            row_fill = fill_mismatch
        else:
            row_fill = fill_missing

        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_num, column=c_idx)
            cell.font = font_data
            cell.fill = row_fill
            cell.border = border_thin
            if c_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [3, 4, 7, 8]:
                cell.font = font_mono
            elif c_idx in [5, 6]:
                cell.alignment = Alignment(horizontal="right", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 11
    ws.column_dimensions["F"].width = 11
    ws.column_dimensions["G"].width = 28
    ws.column_dimensions["H"].width = 28
    ws.column_dimensions["I"].width = 26
    ws.column_dimensions["J"].width = 26
    ws.column_dimensions["K"].width = 38

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    wb.close()
    return output_path
