"""
Model Comparator Engine & PCB Drawing Inspector
Compares two Working Manual PDFs (different series or revisions of a model),
identifies component differences (Added, Removed, Changed Value),
and maps them to physical coordinates on the PCB layout drawings.
"""

import os
import sys
import re
from datetime import datetime

if sys.platform == "win32":
    for p in [r"C:\Program Files\Common Files\microsoft shared\ClickToRun", r"C:\Program Files\Microsoft Office\root\Client"]:
        if os.path.exists(p):
            try:
                os.add_dll_directory(p)
            except Exception:
                pass

import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from bom_extractor import extract_full_bom, extract_metadata
from bom_comparator import parse_locations_string, normalize_part_no, normalize_key

MAX_FILE_SIZE_MB = 10
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024


def clean_filename(filename: str) -> str:
    """
    Strips noise prefix and suffix tags from a file name (e.g. MP, PP, (1), BOM, etc.)
    to isolate the canonical PCB Model and Series.
    Examples:
      'CHP3178AF-1A MP.pdf' -> 'CHP3178AF-1A'
      'CHP3178AF-1A MP (1).pdf' -> 'CHP3178AF-1A'
      'RAD3125AF-2A MP.pdf' -> 'RAD3125AF-2A'
      'BOM_CHP2902DPL-1A (1).xlsx' -> 'CHP2902DPL-1A'
    """
    name = os.path.splitext(os.path.basename(filename))[0]
    # Remove prefix tags
    name = re.sub(r'^(?:BOM|WM|Working[_\s\-]*Manual|Manual|Model[_\s\-]*[AB]?)[_\-\s]+', '', name, flags=re.IGNORECASE)
    # Repeatedly strip trailing copy brackets, stage suffixes, and doc tags
    for _ in range(5):
        prev = name
        # Strip trailing duplicate / copy brackets like (1), [2], Copy, Copy (1)
        name = re.sub(r'[\s_\-]+(?:\(\d+\)|\[\d+\]|Copy(?:\s*\(\d+\))?)$', '', name, flags=re.IGNORECASE)
        # Strip trailing manufacturing stage tags (MP, PP, EVT, DVT, PVT, ES, CS, MASS, PROD, SAMPLE, PILOT, TEST, TP, etc.)
        name = re.sub(r'[\s_\-]+(?:MP|PP|EVT|DVT|PVT|ES|CS|MASS|PROD|SAMPLE|PILOT|TEST|TP|FINAL|NEW|OLD|DRAFT|RELEASED?)$', '', name, flags=re.IGNORECASE)
        # Strip trailing doc type tags (BOM, WM, DRAWING, DRAW, MANUAL, PCB, SCH, LAYOUT, v\d+, rev\d+)
        name = re.sub(r'[\s_\-]+(?:BOM|WM|DRAWING|DRAW|MANUAL|PCB|SCH|LAYOUT|v\d+|rev\d+)$', '', name, flags=re.IGNORECASE)
        name = name.strip(' _-')
        if name == prev:
            break
    return name


def parse_filename_model_series(filename: str) -> tuple[str, str]:
    """
    Extracts (base_model, series) from a filename.
    Examples:
      'CHP3178AF-1A MP.pdf' -> ('CHP3178AF', '1A')
      'CHP3178AF-1B MP.pdf' -> ('CHP3178AF', '1B')
      '3275-2A.pdf' -> ('3275', '2A')
      '3275-2B.pdf' -> ('3275', '2B')
      'CAR3072_RevB.pdf' -> ('CAR3072', 'B')
      'CHA3072AR-2A.pdf' -> ('CHA3072AR', '2A')
    """
    clean = clean_filename(filename)

    # 1. Check for REV / SERIES / VER keyword e.g. CAR3072_RevB
    m_rev = re.search(r'^(.*?)[-_]?(?:REV|SERIES|VER)[-_]?([0-9]*[A-Za-z])$', clean, re.IGNORECASE)
    if m_rev and len(m_rev.group(1)) >= 2:
        return m_rev.group(1).rstrip('-_').upper(), m_rev.group(2).upper()

    # 2. Check for pattern <name>-[0-9]*[A-Za-z] e.g. CHP3178AF-1A, 3275-2A, CHA3072AR-2A, CAR3072-A
    m_code = re.search(r'^(.*?)[-_]([0-9]+[A-Za-z]|[A-Za-z])$', clean)
    if m_code and len(m_code.group(1)) >= 2:
        return m_code.group(1).upper(), m_code.group(2).upper()

    # 3. Pattern <digits><letter> e.g. 3275A, CAR3072A
    m_end = re.search(r'^(.*?\d+)([A-Za-z])$', clean)
    if m_end and len(m_end.group(1)) >= 2:
        return m_end.group(1).upper(), m_end.group(2).upper()

    return clean.upper(), ""


def extract_model_info(pdf_path: str) -> dict:
    """
    Extracts Model identification and Series/Revision code from a Working Manual PDF.
    Combines:
      1. Filename conventions (e.g. CHP3178AF-1A MP.pdf -> Model: CHP3178AF, Series: 1A).
      2. Embedded PDF metadata: PWB Part Code, Process Code, Model/Variant code.
    """
    fname = os.path.splitext(os.path.basename(pdf_path))[0]
    fn_model, fn_series = parse_filename_model_series(fname)

    pwb_code = ""
    process_code = ""
    model_text = ""
    variant_code = ""

    if os.path.exists(pdf_path):
        try:
            doc = fitz.open(pdf_path)
            for pidx in range(min(3, len(doc))):
                txt = doc[pidx].get_text("text")
                meta = extract_metadata(txt)
                if meta.get("pwb_code") and not pwb_code:
                    pwb_code = meta["pwb_code"]
                if meta.get("process_code") and not process_code and meta["process_code"] != "PROGRAM":
                    process_code = meta["process_code"]
                if meta.get("model") and not model_text:
                    model_text = meta["model"].split('\n')[0].strip()
                m_var = re.search(r'MODEL\s*CODE\s*["\']([A-Z0-9]+)["\']', txt, re.IGNORECASE)
                if m_var and not variant_code:
                    variant_code = m_var.group(1).strip()
        except Exception as e:
            print(f"Error extracting metadata from {pdf_path}: {e}")

    # Determine base model and series
    is_generic_fn = bool(re.match(r'^(?:media_\d+|\d{8,}|scan.*|doc.*)$', fname, re.IGNORECASE))

    proc_series = ""
    if process_code:
        m_ps = re.search(r'[-_]([0-9]*[A-Z]+(?:-[A-Z0-9]+)?)$', process_code)
        proc_series = m_ps.group(1) if m_ps else ""

    # Priority for Series:
    # 1. Filename series (if descriptive, e.g. '1A', '1B')
    # 2. Process code series (e.g. '1A' from 'CHP3178AF-1A', '2A' from 'CHA3072AR-2A')
    # 3. Variant code only if it contains letters (e.g. '2A', avoiding numeric chassis codes like '14')
    # 4. Fallback
    if not is_generic_fn and fn_series:
        display_series = fn_series
    elif proc_series:
        display_series = proc_series
    elif variant_code and any(c.isalpha() for c in variant_code):
        display_series = variant_code
    else:
        display_series = proc_series or variant_code or fn_series or ""

    # Priority for Model:
    # Authoritative PWB code if available in PDF, otherwise filename model
    display_model = pwb_code or fn_model

    return {
        "filepath": pdf_path,
        "filename": os.path.basename(pdf_path),
        "display_model": display_model,
        "display_series": display_series,
        "pwb_code": pwb_code,
        "process_code": process_code,
        "fn_model": fn_model,
        "fn_series": fn_series,
        "is_generic_fn": is_generic_fn
    }


def validate_model_pair(pdf_path_a: str, pdf_path_b: str) -> tuple[bool, str, dict, dict]:
    """
    Validates that:
      1. Both files share the SAME base Model.
      2. Both files have DIFFERENT Series / Revisions.
    Returns (is_valid: bool, error_message: str, info_a: dict, info_b: dict).
    """
    if not pdf_path_a or not pdf_path_b:
        return False, "Vui lòng chọn cả 2 file Model A và Model B.", {}, {}

    info_a = extract_model_info(pdf_path_a)
    info_b = extract_model_info(pdf_path_b)

    # 1. Check duplicate identical file
    if os.path.abspath(pdf_path_a) == os.path.abspath(pdf_path_b):
        return False, (
            "Bạn đang chọn 2 file HOÀN TOÀN GIỐNG NHAU!\n\n"
            "Vui lòng chọn 2 bản vẽ có Series khác nhau để so sánh sự sai khác."
        ), info_a, info_b

    # 2. Check Same Model:
    same_model = False
    if info_a["pwb_code"] and info_b["pwb_code"]:
        same_model = (info_a["pwb_code"].upper() == info_b["pwb_code"].upper())
    elif info_a["fn_model"] and info_b["fn_model"] and not info_a["is_generic_fn"] and not info_b["is_generic_fn"]:
        same_model = (info_a["fn_model"].upper() == info_b["fn_model"].upper())
    else:
        same_model = (info_a["display_model"].upper() == info_b["display_model"].upper())

    # Fallback match by core numeric part (e.g. '3178' in QPWBCAF3178 and CHP3178AF)
    if not same_model and info_a["display_model"] and info_b["display_model"]:
        dm_a = info_a["display_model"].upper()
        dm_b = info_b["display_model"].upper()
        nums_a = re.findall(r'\d{3,}', dm_a)
        nums_b = re.findall(r'\d{3,}', dm_b)
        if nums_a and nums_b and set(nums_a) == set(nums_b):
            same_model = True

    if not same_model:
        m_a = info_a["display_model"]
        m_b = info_b["display_model"]
        return False, (
            f"Hai bản vẽ KHÔNG CÙNG MODEL!\n\n"
            f"• File A: Model '{m_a}'\n"
            f"• File B: Model '{m_b}'\n\n"
            f"Yêu cầu: Hai bản vẽ phải CÙNG MODEL nhưng KHÁC SERIES (ví dụ: cùng model 3275 nhưng là series 1A và 1B)."
        ), info_a, info_b

    # 3. Check Different Series:
    ser_a = info_a["display_series"].upper()
    ser_b = info_b["display_series"].upper()

    if ser_a and ser_b and ser_a == ser_b:
        return False, (
            f"Hai bản vẽ TRÙNG SERIES ('{ser_a}')!\n\n"
            f"• File A: {info_a['filename']} (Series: {ser_a})\n"
            f"• File B: {info_b['filename']} (Series: {ser_b})\n\n"
            f"Yêu cầu: Hai bản vẽ phải KHÁC SERIES (ví dụ: Series 1A và Series 1B) để phát hiện sự sai khác."
        ), info_a, info_b

    return True, "", info_a, info_b


def normalize_process_stage(proc_name: str) -> str:
    """Normalizes various process codes across series into common manufacturing stages."""
    if not proc_name:
        return "SMT"
    p = str(proc_name).upper().strip()
    if any(k in p for k in ["CHA", "REF-A", "REF", "SMT", "CHIP"]):
        return "SMT"
    if any(k in p for k in ["RAD", "RD", "AXI"]):
        return "RADIAL"
    if any(k in p for k in ["SCP", "DIP", "MANUAL", "HAND"]):
        return "DIP"
    return p.split("-")[0] if "-" in p else p


def find_drawing_pages(pdf_path: str) -> list[dict]:
    """
    Identifies which pages in the PDF are PCB layout / assembly drawing pages.
    Accurately classifies SMT, RADIAL/AXI, and DIP/MANUAL manufacturing stages.
    """
    doc = fitz.open(pdf_path)
    drawing_pages = []
    
    for idx, page in enumerate(doc):
        txt = page.get_text("text")
        lines = [l.strip() for l in txt.split("\n") if l.strip()]
        
        proc_code = ""
        for i, l in enumerate(lines):
            if "PROCESS CODE" in l:
                if len(l.split("PROCESS CODE")) > 1 and l.split("PROCESS CODE")[1].strip():
                    proc_code = l.split("PROCESS CODE")[1].strip()
                elif i + 1 < len(lines):
                    proc_code = lines[i + 1]
                break

        stage = ""
        is_drawing = False

        # Classify stage primarily by process code
        p_up = proc_code.upper()
        if any(k in p_up for k in ["CHA", "REF", "SMT", "CHIP"]):
            stage = "SMT"
        elif any(k in p_up for k in ["RAD", "RD", "AXI"]):
            stage = "RADIAL"
        elif any(k in p_up for k in ["SCP", "DIP", "MANUAL", "HAND"]):
            stage = "DIP"

        # Secondary check by page text keywords
        txt_up = txt.upper()
        if not stage:
            if "RAD3072" in txt_up or "AXI/RAD" in txt_up or "AR RD" in txt_up:
                stage = "RADIAL"
            elif "SCP3072" in txt_up or "AR DIP" in txt_up or "DIP A-SIDE" in txt_up or "DIP B" in txt_up:
                stage = "DIP"
            elif "CHA3072" in txt_up or "REF-A" in txt_up or "REFLOW" in txt_up:
                stage = "SMT"

        # Determine if page contains drawing content
        if "PAGE NO." in txt and "2/2" in txt:
            is_drawing = True
            if not stage:
                stage = "SMT"
        elif "FLOW DIRECTION" in txt or "AXI/RAD PCB" in txt or "CHIP PCB FLOW DIRECTION" in txt:
            is_drawing = True
        elif "SIDE A VIEW" in txt or "SIDE B VIEW" in txt or "AOI INSPECTION" in txt or "VISUAL INSPECTION" in txt:
            is_drawing = True
        elif "STAMPING AREA" in txt:
            is_drawing = True

        if is_drawing:
            if not stage:
                stage = "SMT"
            drawing_pages.append({
                "page_idx": idx,
                "stage": stage,
                "process_code": proc_code,
                "rotation": page.rotation,
                "rect": (page.rect.x0, page.rect.y0, page.rect.x1, page.rect.y1)
            })
            
    doc.close()
    return drawing_pages


def get_drawing_pages_catalog(pdf_path_a: str, pdf_path_b: str) -> list[dict]:
    """
    Builds a catalog of available drawing pages and stages for user dropdown selection.
    """
    drawings_a = find_drawing_pages(pdf_path_a) if (pdf_path_a and os.path.exists(pdf_path_a)) else []
    drawings_b = find_drawing_pages(pdf_path_b) if (pdf_path_b and os.path.exists(pdf_path_b)) else []

    catalog = [
        {
            "id": "ALL",
            "stage": "ALL",
            "label": "🌐 Tất Cả Công Đoạn (Tự Động Theo LK)",
            "page_idx_a": drawings_a[0]["page_idx"] if drawings_a else 1,
            "page_idx_b": drawings_b[0]["page_idx"] if drawings_b else 1
        }
    ]

    stages_order = [
        ("SMT", "⚡", "SMT / CHIP"),
        ("RADIAL", "🔌", "RADIAL / AXI"),
        ("DIP", "🛠️", "DIP / MANUAL")
    ]

    for st_code, icon, st_name in stages_order:
        da = next((d for d in drawings_a if d["stage"] == st_code), None)
        db = next((d for d in drawings_b if d["stage"] == st_code), None)
        if da or db:
            pa = da["page_idx"] if da else (db["page_idx"] if db else 1)
            pb = db["page_idx"] if db else (da["page_idx"] if da else 1)
            p_code = (da.get("process_code") if da else "") or (db.get("process_code") if db else "") or st_code
            page_num = (pa + 1) if da else (pb + 1)
            catalog.append({
                "id": st_code,
                "stage": st_code,
                "label": f"{icon} Trang {page_num}: {st_name} ({p_code})",
                "page_idx_a": pa,
                "page_idx_b": pb
            })

    existing_pages_a = {c["page_idx_a"] for c in catalog if c["id"] != "ALL"}
    for d in drawings_a:
        if d["page_idx"] not in existing_pages_a:
            catalog.append({
                "id": f"PAGE_{d['page_idx']}",
                "stage": d["stage"],
                "label": f"📄 Trang {d['page_idx']+1}: {d['stage']} ({d.get('process_code')})",
                "page_idx_a": d["page_idx"],
                "page_idx_b": d["page_idx"]
            })

    return catalog


def locate_ref_on_drawing(doc: fitz.Document, page_idx: int, ref_des: str) -> list[tuple[float, float, float, float]]:
    """
    Locates reference designator text bounding boxes on an unrotated coordinate space.
    Returns list of (x0, y0, x1, y1) in 72 DPI PDF points.
    """
    if page_idx < 0 or page_idx >= len(doc):
        return []
    page = doc[page_idx]
    
    # Save original rotation and set to 0 for consistent bounding coordinates
    orig_rot = page.rotation
    page.set_rotation(0)
    
    rects = page.search_for(ref_des)
    
    # Filter exact matches to avoid partial hits like C1 matching C10, C11
    exact_rects = []
    ref_upper = ref_des.upper().strip()
    words = page.get_text("words") # (x0, y0, x1, y1, word, block_no, line_no, word_no)
    for w in words:
        if w[4].upper().strip() == ref_upper:
            exact_rects.append((w[0], w[1], w[2], w[3]))
            
    page.set_rotation(orig_rot)
    
    if exact_rects:
        return exact_rects
    return [(r.x0, r.y0, r.x1, r.y1) for r in rects]


def compare_model_manuals(pdf_path_a: str, pdf_path_b: str) -> dict:
    """
    Compares two Working Manual PDFs and returns structured diff data.
    Categories:
      - ADDED: Present in Model B but not Model A
      - REMOVED: Present in Model A but not Model B
      - CHANGED: Present in both, but Part Number or Rating/Spec changed
      - MATCH: Present in both with identical Part Number & Rating
    """
    for p in [pdf_path_a, pdf_path_b]:
        if not os.path.exists(p):
            raise FileNotFoundError(f"File not found: {p}")
        if os.path.getsize(p) > MAX_FILE_SIZE_BYTES:
            raise ValueError(f"File '{os.path.basename(p)}' vượt quá giới hạn {MAX_FILE_SIZE_MB}MB.")

    # Validate that both files belong to the same model but have different series
    is_valid, err_msg, info_a, info_b = validate_model_pair(pdf_path_a, pdf_path_b)
    if not is_valid:
        raise ValueError(err_msg)

    # Extract BOM items using core extractor
    meta_a, flat_a, _ = extract_full_bom(pdf_path_a)
    meta_b, flat_b, _ = extract_full_bom(pdf_path_b)

    # Detect drawing pages
    drawings_a = find_drawing_pages(pdf_path_a)
    drawings_b = find_drawing_pages(pdf_path_b)

    # Index items by (stage, ref_des)
    map_a: dict[tuple[str, str], dict] = {}
    for r in flat_a:
        stage = normalize_process_stage(r.get("Process", ""))
        locs = parse_locations_string(r.get("Remarks_Locations", ""))
        for loc in locs:
            map_a[(stage, loc)] = r

    map_b: dict[tuple[str, str], dict] = {}
    for r in flat_b:
        stage = normalize_process_stage(r.get("Process", ""))
        locs = parse_locations_string(r.get("Remarks_Locations", ""))
        for loc in locs:
            map_b[(stage, loc)] = r

    all_keys = sorted(list(set(map_a.keys()) | set(map_b.keys())))

    diff_items = []
    match_items = []

    for stage, loc in all_keys:
        in_a = (stage, loc) in map_a
        in_b = (stage, loc) in map_b

        if in_a and not in_b:
            item_a = map_a[(stage, loc)]
            diff_items.append({
                "status": "REMOVED",
                "stage": stage,
                "ref_des": loc,
                "part_a": item_a.get("Part_No", ""),
                "rating_a": item_a.get("Rating_Spec", ""),
                "qty_a": 1,
                "part_b": "---",
                "rating_b": "DNP / Bỏ trống",
                "qty_b": 0,
                "note": "Chỉ có ở Model A (Model B bỏ không lắp)"
            })
        elif in_b and not in_a:
            item_b = map_b[(stage, loc)]
            diff_items.append({
                "status": "ADDED",
                "stage": stage,
                "ref_des": loc,
                "part_a": "---",
                "rating_a": "DNP / Bỏ trống",
                "qty_a": 0,
                "part_b": item_b.get("Part_No", ""),
                "rating_b": item_b.get("Rating_Spec", ""),
                "qty_b": 1,
                "note": "Chỉ có ở Model B (Model B lắp mới)"
            })
        else:
            item_a = map_a[(stage, loc)]
            item_b = map_b[(stage, loc)]
            pn_a = item_a.get("Part_No", "")
            pn_g = item_b.get("Part_No", "")
            rat_a = item_a.get("Rating_Spec", "")
            rat_g = item_b.get("Rating_Spec", "")

            # Check if part numbers or ratings differ
            if normalize_key(pn_a) != normalize_key(pn_g) or rat_a.strip() != rat_g.strip():
                diff_items.append({
                    "status": "CHANGED",
                    "stage": stage,
                    "ref_des": loc,
                    "part_a": pn_a,
                    "rating_a": rat_a,
                    "qty_a": 1,
                    "part_b": pn_g,
                    "rating_b": rat_g,
                    "qty_b": 1,
                    "note": f"Đổi linh kiện: {pn_a} -> {pn_g}"
                })
            else:
                match_items.append({
                    "status": "MATCH",
                    "stage": stage,
                    "ref_des": loc,
                    "part_a": pn_a,
                    "rating_a": rat_a,
                    "qty_a": 1,
                    "part_b": pn_g,
                    "rating_b": rat_g,
                    "qty_b": 1,
                    "note": "Trùng khớp hoàn toàn"
                })

    added_count = sum(1 for d in diff_items if d["status"] == "ADDED")
    removed_count = sum(1 for d in diff_items if d["status"] == "REMOVED")
    changed_count = sum(1 for d in diff_items if d["status"] == "CHANGED")

    result = {
        "model_a": {
            "path": pdf_path_a,
            "filename": os.path.basename(pdf_path_a),
            "display_model": info_a.get("display_model", ""),
            "display_series": info_a.get("display_series", ""),
            "pwb_code": meta_a.get("pwb_code", ""),
            "models": meta_a.get("models", []),
            "processes": meta_a.get("processes", []),
            "total_items": len(flat_a),
            "drawings": drawings_a
        },
        "model_b": {
            "path": pdf_path_b,
            "filename": os.path.basename(pdf_path_b),
            "display_model": info_b.get("display_model", ""),
            "display_series": info_b.get("display_series", ""),
            "pwb_code": meta_b.get("pwb_code", ""),
            "models": meta_b.get("models", []),
            "processes": meta_b.get("processes", []),
            "total_items": len(flat_b),
            "drawings": drawings_b
        },
        "summary": {
            "total_positions": len(all_keys),
            "match_count": len(match_items),
            "diff_count": len(diff_items),
            "added_count": added_count,
            "removed_count": removed_count,
            "changed_count": changed_count,
            "is_identical": (len(diff_items) == 0)
        },
        "diff_items": diff_items,
        "match_items": match_items,
        "all_items": diff_items + match_items,
        "drawing_catalog": get_drawing_pages_catalog(pdf_path_a, pdf_path_b)
    }
    return result


def render_component_spotlight(
    pdf_path: str,
    page_idx: int,
    target_ref: str,
    status: str = "ADDED",
    crop_size: tuple[int, int] = (420, 420),
    zoom: float = 2.0
) -> Image.Image:
    """
    Renders a zoomed-in crop of the PCB drawing centered on target_ref,
    with a highlighted box, indicator badge, and clear label.
    """
    doc = fitz.open(pdf_path)
    if page_idx < 0 or page_idx >= len(doc):
        doc.close()
        # Return blank placeholder image
        img = Image.new("RGB", crop_size, "#1E293B")
        draw = ImageDraw.Draw(img)
        draw.text((crop_size[0]//2 - 60, crop_size[1]//2 - 10), "Không có bản vẽ", fill="#94A3B8")
        return img

    page = doc[page_idx]
    orig_rot = page.rotation
    page.set_rotation(0)
    
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat)
    full_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

    rects = locate_ref_on_drawing(doc, page_idx, target_ref)
    page.set_rotation(orig_rot)
    doc.close()

    if not rects:
        # If text not found on drawing page, return center of page crop with warning
        cx, cy = full_img.width // 2, full_img.height // 2
    else:
        rx0, ry0, rx1, ry1 = rects[0]
        # Map 72 DPI to zoom pixels
        cx = int(((rx0 + rx1) / 2.0) * zoom)
        cy = int(((ry0 + ry1) / 2.0) * zoom)

    # Determine status color
    clr_map = {
        "ADDED": "#10B981",    # Emerald Green
        "REMOVED": "#EF4444",  # Coral Red
        "CHANGED": "#F59E0B",  # Amber
        "MATCH": "#3B82F6"     # Blue
    }
    box_color = clr_map.get(status, "#EF4444")

    # Crop around (cx, cy)
    w, h = crop_size
    x0 = max(0, min(full_img.width - w, cx - w // 2))
    y0 = max(0, min(full_img.height - h, cy - h // 2))
    x1 = x0 + w
    y1 = y0 + h

    cropped = full_img.crop([x0, y0, x1, y1])
    draw = ImageDraw.Draw(cropped)

    # Draw highlight box on target rect if inside crop
    if rects:
        rx0, ry0, rx1, ry1 = rects[0]
        px0 = int(rx0 * zoom) - x0 - 8
        py0 = int(ry0 * zoom) - y0 - 8
        px1 = int(rx1 * zoom) - x0 + 8
        py1 = int(ry1 * zoom) - y0 + 8

        # Draw outer glowing box + inner outline
        draw.rectangle([px0 - 2, py0 - 2, px1 + 2, py1 + 2], outline=box_color, width=4)
        draw.rectangle([px0 + 2, py0 + 2, px1 - 2, py1 - 2], outline="#FFFFFF", width=1)

        # Draw callout badge banner at top-left of crop
        badge_txt = f"📍 {target_ref} ({status})"
        draw.rectangle([8, 8, 170, 32], fill="#0F172A")
        draw.rectangle([8, 8, 170, 32], outline=box_color, width=2)
        draw.text((16, 12), badge_txt, fill=box_color)

    return cropped


def render_full_drawing_view(
    pdf_path: str,
    page_idx: int,
    active_ref: str = None,
    all_diff_refs: list[dict] = None,
    zoom: float = 1.2
) -> Image.Image:
    """
    Renders the full PCB drawing page with highlights for all diff locations,
    and a prominent spotlight on active_ref.
    """
    doc = fitz.open(pdf_path)
    if page_idx < 0 or page_idx >= len(doc):
        doc.close()
        return Image.new("RGB", (800, 600), "#1E293B")

    page = doc[page_idx]
    orig_rot = page.rotation
    page.set_rotation(0)

    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat)
    full_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

    draw = ImageDraw.Draw(full_img)

    # Highlight all diff refs subtly
    if all_diff_refs:
        for item in all_diff_refs:
            ref = item.get("ref_des", "")
            st = item.get("status", "CHANGED")
            clr = "#10B981" if st == "ADDED" else ("#EF4444" if st == "REMOVED" else "#F59E0B")
            r_list = locate_ref_on_drawing(doc, page_idx, ref)
            for r in r_list:
                px0 = int(r[0] * zoom) - 4
                py0 = int(r[1] * zoom) - 4
                px1 = int(r[2] * zoom) + 4
                py1 = int(r[3] * zoom) + 4
                draw.rectangle([px0, py0, px1, py1], outline=clr, width=2)

    # Highlight active ref prominently
    if active_ref:
        act_rects = locate_ref_on_drawing(doc, page_idx, active_ref)
        for r in act_rects:
            px0 = int(r[0] * zoom) - 8
            py0 = int(r[1] * zoom) - 8
            px1 = int(r[2] * zoom) + 8
            py1 = int(r[3] * zoom) + 8
            draw.rectangle([px0, py0, px1, py1], outline="#DC2626", width=5)
            draw.rectangle([px0 - 2, py0 - 2, px1 + 2, py1 + 2], outline="#FDE047", width=2)

    page.set_rotation(orig_rot)
    doc.close()
    return full_img


_CURTAIN_BASE_CACHE: dict = {}

def clear_curtain_cache():
    """Clears cached base drawing images."""
    global _CURTAIN_BASE_CACHE
    _CURTAIN_BASE_CACHE.clear()

def get_annotated_base_images(
    pdf_path_a: str,
    pdf_path_b: str,
    page_idx_a: int,
    page_idx_b: int,
    zoom: float = 1.0,
    diff_items: list[dict] = None,
    active_ref: str = None,
    rotation: int = 0
) -> tuple[Image.Image, Image.Image, tuple[int, int] | None, int, int]:
    """
    Renders and caches annotated base images for Model A and Model B with rotation support.
    Because this is cached, curtain dragging does NOT re-rasterize PDFs or search text!
    """
    global _CURTAIN_BASE_CACHE
    diff_sig = tuple((d.get("ref_des"), d.get("status")) for d in diff_items) if diff_items else ()
    rot_deg = int(rotation) % 360
    cache_key = (
        os.path.abspath(pdf_path_a) if pdf_path_a else "",
        os.path.abspath(pdf_path_b) if pdf_path_b else "",
        page_idx_a,
        page_idx_b,
        round(float(zoom), 3),
        active_ref or "",
        diff_sig,
        rot_deg
    )
    if cache_key in _CURTAIN_BASE_CACHE:
        return _CURTAIN_BASE_CACHE[cache_key]

    doc_a = fitz.open(pdf_path_a) if (pdf_path_a and os.path.exists(pdf_path_a)) else None
    doc_b = fitz.open(pdf_path_b) if (pdf_path_b and os.path.exists(pdf_path_b)) else None

    if not doc_a and not doc_b:
        blank = Image.new("RGB", (800, 600), "#0B0F1A")
        return blank, blank, None, 800, 600

    # Clamp page indices safely
    if doc_a and (page_idx_a < 0 or page_idx_a >= len(doc_a)):
        page_idx_a = min(len(doc_a) - 1, max(0, page_idx_a))
    if doc_b and (page_idx_b < 0 or page_idx_b >= len(doc_b)):
        page_idx_b = min(len(doc_b) - 1, max(0, page_idx_b))

    page_a = doc_a[page_idx_a] if doc_a else None
    page_b = doc_b[page_idx_b] if doc_b else None

    orig_rot_a = 0
    orig_rot_b = 0
    if page_a:
        orig_rot_a = page_a.rotation
        page_a.set_rotation(0)
    if page_b:
        orig_rot_b = page_b.rotation
        page_b.set_rotation(0)

    mat = fitz.Matrix(zoom, zoom)

    img_a = None
    if page_a:
        pix_a = page_a.get_pixmap(matrix=mat)
        img_a = Image.frombytes("RGB", [pix_a.width, pix_a.height], pix_a.samples)

    img_b = None
    if page_b:
        pix_b = page_b.get_pixmap(matrix=mat)
        img_b = Image.frombytes("RGB", [pix_b.width, pix_b.height], pix_b.samples)

    if not img_a and img_b:
        img_a = img_b.copy()
    if not img_b and img_a:
        img_b = img_a.copy()

    w = min(img_a.width, img_b.width)
    h = min(img_a.height, img_b.height)

    if img_a.size != (w, h):
        img_a = img_a.crop((0, 0, w, h))
    if img_b.size != (w, h):
        img_b = img_b.crop((0, 0, w, h))

    orig_w, orig_h = w, h

    # Apply rotation if specified (0, 90, 180, 270 degrees)
    if rot_deg == 90:
        img_a = img_a.transpose(Image.ROTATE_270)
        img_b = img_b.transpose(Image.ROTATE_270)
        w, h = orig_h, orig_w
    elif rot_deg == 180:
        img_a = img_a.transpose(Image.ROTATE_180)
        img_b = img_b.transpose(Image.ROTATE_180)
    elif rot_deg == 270:
        img_a = img_a.transpose(Image.ROTATE_90)
        img_b = img_b.transpose(Image.ROTATE_90)
        w, h = orig_h, orig_w

    def _rot_box(box, deg, bw, bh):
        bx0, by0, bx1, by1 = box
        if deg == 90:
            return [bh - 1 - by1, bx0, bh - 1 - by0, bx1]
        elif deg == 180:
            return [bw - 1 - bx1, bh - 1 - by1, bw - 1 - bx0, bh - 1 - by0]
        elif deg == 270:
            return [by0, bw - 1 - bx1, by1, bw - 1 - bx0]
        return [bx0, by0, bx1, by1]

    draw_a = ImageDraw.Draw(img_a)
    draw_b = ImageDraw.Draw(img_b)

    active_center = None

    # Draw subtle, thin bounding boxes (width=1) for differences on this page
    if diff_items:
        for item in diff_items:
            ref = item.get("ref_des", "")
            st = item.get("status", "CHANGED")
            clr = "#10B981" if st == "ADDED" else ("#EF4444" if st == "REMOVED" else "#F59E0B")

            if st in ("ADDED", "CHANGED") and doc_b:
                r_list_b = locate_ref_on_drawing(doc_b, page_idx_b, ref)
                for r in r_list_b:
                    raw_b = [int(r[0] * zoom) - 2, int(r[1] * zoom) - 2, int(r[2] * zoom) + 2, int(r[3] * zoom) + 2]
                    rb = _rot_box(raw_b, rot_deg, orig_w, orig_h)
                    px0 = max(0, min(w, min(rb[0], rb[2])))
                    py0 = max(0, min(h, min(rb[1], rb[3])))
                    px1 = max(0, min(w, max(rb[0], rb[2])))
                    py1 = max(0, min(h, max(rb[1], rb[3])))
                    draw_b.rectangle([px0, py0, px1, py1], outline=clr, width=1)

            if st in ("REMOVED", "CHANGED") and doc_a:
                r_list_a = locate_ref_on_drawing(doc_a, page_idx_a, ref)
                for r in r_list_a:
                    raw_a = [int(r[0] * zoom) - 2, int(r[1] * zoom) - 2, int(r[2] * zoom) + 2, int(r[3] * zoom) + 2]
                    ra = _rot_box(raw_a, rot_deg, orig_w, orig_h)
                    px0 = max(0, min(w, min(ra[0], ra[2])))
                    py0 = max(0, min(h, min(ra[1], ra[3])))
                    px1 = max(0, min(w, max(ra[0], ra[2])))
                    py1 = max(0, min(h, max(ra[1], ra[3])))
                    draw_a.rectangle([px0, py0, px1, py1], outline=clr, width=1)

    # Draw refined, thin spotlight for active_ref (width=2) with upright label
    if active_ref:
        r_act_b = locate_ref_on_drawing(doc_b, page_idx_b, active_ref) if doc_b else []
        r_act_a = locate_ref_on_drawing(doc_a, page_idx_a, active_ref) if doc_a else []
        r_act = r_act_b or r_act_a

        if r_act:
            rx0, ry0, rx1, ry1 = r_act[0]
            raw_act = [int(rx0 * zoom) - 4, int(ry0 * zoom) - 4, int(rx1 * zoom) + 4, int(ry1 * zoom) + 4]
            r_box = _rot_box(raw_act, rot_deg, orig_w, orig_h)
            px0 = max(0, min(w, min(r_box[0], r_box[2])))
            py0 = max(0, min(h, min(r_box[1], r_box[3])))
            px1 = max(0, min(w, max(r_box[0], r_box[2])))
            py1 = max(0, min(h, max(r_box[1], r_box[3])))

            cx = (px0 + px1) // 2
            cy = (py0 + py1) // 2
            active_center = (cx, cy)

            for draw_obj in [draw_a, draw_b]:
                # Clean thin 2px primary border + 1px subtle yellow inner border
                draw_obj.rectangle([px0, py0, px1, py1], outline="#DC2626", width=2)
                draw_obj.rectangle([px0 - 1, py0 - 1, px1 + 1, py1 + 1], outline="#FDE047", width=1)
                # Upright horizontal callout label
                draw_obj.rectangle([px0, max(0, py0 - 18), px0 + 84, py0], fill="#0F172A")
                draw_obj.rectangle([px0, max(0, py0 - 18), px0 + 84, py0], outline="#FDE047", width=1)
                draw_obj.text((px0 + 4, max(0, py0 - 15)), f"📍 {active_ref}", fill="#FDE047")

    if page_a:
        page_a.set_rotation(orig_rot_a)
    if page_b:
        page_b.set_rotation(orig_rot_b)

    if doc_a:
        doc_a.close()
    if doc_b:
        doc_b.close()

    # Cache up to 8 base drawing pairs
    if len(_CURTAIN_BASE_CACHE) >= 8:
        _CURTAIN_BASE_CACHE.pop(next(iter(_CURTAIN_BASE_CACHE)))
    _CURTAIN_BASE_CACHE[cache_key] = (img_a, img_b, active_center, w, h)
    return img_a, img_b, active_center, w, h


def render_curtain_drawing_pair(
    pdf_path_a: str,
    pdf_path_b: str,
    page_idx_a: int,
    page_idx_b: int,
    zoom: float = 1.0,
    split_ratio: float = 0.5,
    mode: str = "curtain",  # "curtain", "a", "b"
    diff_items: list[dict] = None,
    active_ref: str = None,
    rotation: int = 0
) -> tuple[Image.Image, dict]:
    """
    Renders a paired drawing comparison view with curtain wipe slider effect,
    component spotlights, difference highlights across the drawing, and rotation.
    Leverages cached base images for ultra-smooth 60+ FPS curtain dragging.
    """
    img_a, img_b, active_center, w, h = get_annotated_base_images(
        pdf_path_a, pdf_path_b, page_idx_a, page_idx_b, zoom, diff_items, active_ref, rotation
    )

    if mode == "a":
        res_img = img_a.copy()
        split_x = None
    elif mode == "b":
        res_img = img_b.copy()
        split_x = None
    else:  # "curtain"
        split_ratio = max(0.0, min(1.0, split_ratio))
        split_x = max(0, min(w, int(w * split_ratio)))
        res_img = img_b.copy()
        if split_x > 0:
            res_img.paste(img_a.crop((0, 0, split_x, h)), (0, 0))

        draw_c = ImageDraw.Draw(res_img)
        # Clean thin cyan curtain divider line
        draw_c.line([(split_x, 0), (split_x, h)], fill="#22D3EE", width=2)

        # Top Model Badges (Safely bounded within canvas margins)
        b_y = 12
        if split_x > 30:
            a_x2 = min(split_x - 6, w - 6)
            a_x1 = max(4, a_x2 - 110)
            if a_x2 > a_x1:
                draw_c.rectangle([a_x1, b_y, a_x2, b_y + 24], fill="#0F172A")
                draw_c.rectangle([a_x1, b_y, a_x2, b_y + 24], outline="#22D3EE", width=1)
                draw_c.text((a_x1 + 6, b_y + 5), "📄 Model A (Gốc)", fill="#38BDF8")

        if split_x < w - 30:
            b_x1 = max(4, split_x + 6)
            b_x2 = min(w - 4, b_x1 + 110)
            if b_x2 > b_x1:
                draw_c.rectangle([b_x1, b_y, b_x2, b_y + 24], fill="#0F172A")
                draw_c.rectangle([b_x1, b_y, b_x2, b_y + 24], outline="#10B981", width=1)
                draw_c.text((b_x1 + 6, b_y + 5), "📄 Model B (Mới)", fill="#34D399")

        # Center draggable handle button (Safely clamped)
        my = h // 2
        hx = max(14, min(w - 14, split_x))
        draw_c.ellipse([hx - 14, my - 14, hx + 14, my + 14], fill="#0891B2", outline="#FFFFFF", width=2)
        draw_c.text((hx - 9, my - 7), "◂||▸", fill="#FFFFFF")

    meta = {
        "width": w,
        "height": h,
        "split_x": split_x,
        "split_ratio": split_ratio,
        "active_center": active_center,
        "mode": mode,
        "rotation": rotation % 360
    }
    return res_img, meta


def export_model_comparison_excel(comp_result: dict, output_path: str) -> str:
    """
    Exports a professional ECN Delta Report Excel file comparing the two models.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model_Diff_Summary"
    ws.views.sheetView[0].showGridLines = True

    # Styles
    f_title = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    f_sub = Font(name="Segoe UI", size=10, italic=True, color="475569")
    f_card_hdr = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    f_hdr = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    f_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    f_cell = Font(name="Segoe UI", size=10, color="0F172A")

    fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_blue_hdr = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    fill_added = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")   # Green
    fill_removed = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Red
    fill_changed = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Amber
    fill_match = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Title Banner
    ws.merge_cells("B2:H2")
    ws["B2"] = "BÁO CÁO SO SÁNH KHÁC BIỆT 2 MODEL (ECN DELTA REPORT)"
    ws["B2"].font = f_title
    ws["B2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 30

    ws.merge_cells("B3:H3")
    ws["B3"] = f"Thời gian tạo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}  |  Bản quyền: Lê Văn Vững - IPQC"
    ws["B3"].font = f_sub
    ws.row_dimensions[3].height = 20

    # Model Overview Cards
    mod_a = comp_result.get("model_a", {})
    mod_b = comp_result.get("model_b", {})
    sum_data = comp_result.get("summary", {})

    ws["B5"] = "THÔNG TIN MODEL A (CŨ / GỐC)"
    ws["B5"].font = f_card_hdr
    ws["B5"].fill = fill_navy
    ws["C5"] = mod_a.get("filename", "")
    ws["C5"].font = f_bold

    ws["B6"] = "Processes: " + ", ".join(mod_a.get("processes", []))
    ws["B6"].font = f_cell
    ws["C6"] = f"Tổng số linh kiện: {mod_a.get('total_items', 0)}"
    ws["C6"].font = f_bold

    ws["E5"] = "THÔNG TIN MODEL B (MỚI / ĐỔI)"
    ws["E5"].font = f_card_hdr
    ws["E5"].fill = fill_navy
    ws["F5"] = mod_b.get("filename", "")
    ws["F5"].font = f_bold

    ws["E6"] = "Processes: " + ", ".join(mod_b.get("processes", []))
    ws["E6"].font = f_cell
    ws["F6"] = f"Tổng số linh kiện: {mod_b.get('total_items', 0)}"
    ws["F6"].font = f_bold

    # Summary Statistics Row
    ws.row_dimensions[8].height = 25
    headers_sum = [
        ("B8", f"TỔNG VỊ TRÍ: {sum_data.get('total_positions', 0)}", "2563EB"),
        ("C8", f"TRÙNG KHỚP: {sum_data.get('match_count', 0)}", "059669"),
        ("D8", f"LẮP THÊM: {sum_data.get('added_count', 0)}", "10B981"),
        ("E8", f"BỎ BỚT: {sum_data.get('removed_count', 0)}", "EF4444"),
        ("F8", f"ĐỔI TRỊ SỐ: {sum_data.get('changed_count', 0)}", "D97706"),
    ]
    for cell_id, text, color in headers_sum:
        ws[cell_id] = text
        ws[cell_id].font = f_card_hdr
        ws[cell_id].fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws[cell_id].alignment = Alignment(horizontal="center", vertical="center")

    # Table Header Row
    headers = [
        ("STT", 6),
        ("CÔNG ĐOẠN", 14),
        ("VỊ TRÍ (REF DES)", 18),
        ("TRẠNG THÁI", 16),
        ("MODEL A (PART NO)", 26),
        ("MODEL A (RATING/SPEC)", 24),
        ("MODEL B (PART NO)", 26),
        ("MODEL B (RATING/SPEC)", 24),
        ("GHI CHÚ KỸ THUẬT", 35)
    ]

    r_idx = 10
    ws.row_dimensions[r_idx].height = 26
    for col_idx, (h_title, w) in enumerate(headers, start=2):
        cell = ws.cell(row=r_idx, column=col_idx, value=h_title)
        cell.font = f_hdr
        cell.fill = fill_blue_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = w

    # Table Data Rows: Diff items first, then matches
    all_rows = comp_result.get("diff_items", []) + comp_result.get("match_items", [])
    for idx, item in enumerate(all_rows, start=1):
        r_idx += 1
        st = item.get("status", "MATCH")
        if st == "ADDED":
            row_fill = fill_added
            st_text = "🟢 LẮP THÊM"
        elif st == "REMOVED":
            row_fill = fill_removed
            st_text = "🔴 BỎ BỚT"
        elif st == "CHANGED":
            row_fill = fill_changed
            st_text = "🟡 ĐỔI TRỊ SỐ"
        else:
            row_fill = fill_match
            st_text = "⚪ TRÙNG KHỚP"

        row_vals = [
            idx,
            item.get("stage", ""),
            item.get("ref_des", ""),
            st_text,
            item.get("part_a", ""),
            item.get("rating_a", ""),
            item.get("part_b", ""),
            item.get("rating_b", ""),
            item.get("note", "")
        ]

        ws.row_dimensions[r_idx].height = 20
        for col_idx, val in enumerate(row_vals, start=2):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = f_cell
            cell.fill = row_fill
            cell.border = thin_border
            align = "center" if col_idx in (2, 3, 4, 5) else "left"
            cell.alignment = Alignment(horizontal=align, vertical="center")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path
